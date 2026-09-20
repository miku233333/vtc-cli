from pathlib import Path

from docx import Document
from pptx import Presentation

from vtc.cli import main
from vtc.moodle.extract import extract_file
from vtc.moodle.sync import _attach_extract


def _write_pdf(path: Path, text: str) -> None:
    safe = "".join(ch if ch.isascii() and ch not in "()\\" else " " for ch in text)
    stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET\n".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    body = b"%PDF-1.4\n"
    offsets: list[int] = []
    for item in objects:
        offsets.append(len(body))
        n = len(offsets)
        body += f"{n} 0 obj\n".encode("ascii") + item + b"\nendobj\n"
    xref_pos = len(body)
    xref = [b"xref\n0 6\n0000000000 65535 f \n"]
    for offset in offsets:
        xref.append(f"{offset:010d} 00000 n \n".encode("ascii"))
    path.write_bytes(
        body
        + b"".join(xref)
        + b"trailer\n<< /Size 6 /Root 1 0 R >>\n"
        + f"startxref\n{xref_pos}\n".encode("ascii")
        + b"%%EOF\n"
    )


def test_extract_pdf_docx_pptx(tmp_path):
    pdf = tmp_path / "lecture.pdf"
    docx = tmp_path / "notes.docx"
    pptx = tmp_path / "slides.pptx"
    _write_pdf(pdf, "Hello PDF from Moodle")
    document = Document()
    document.add_paragraph("Hello Word from Moodle")
    document.save(docx)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "Hello PPT from Moodle"
    deck.save(pptx)

    pdf_text = extract_file(pdf)
    docx_text = extract_file(docx)
    pptx_text = extract_file(pptx)
    assert pdf_text.status == "ok"
    assert "Hello PDF from Moodle" in (pdf_text.text or "")
    assert docx_text.status == "ok"
    assert "Hello Word from Moodle" in (docx_text.text or "")
    assert pptx_text.status == "ok"
    assert "Hello PPT from Moodle" in (pptx_text.text or "")


def test_extract_legacy_ppt_is_unverified_without_textutil(tmp_path, monkeypatch):
    path = tmp_path / "old.ppt"
    path.write_bytes(b"not a real ppt")
    monkeypatch.setattr("vtc.moodle.extract._extract_textutil", lambda *_args, **_kwargs: (None, None))
    result = extract_file(path)
    assert result.status == "unverified"
    assert result.error == "legacy_office_unreadable"


def test_sync_attach_extract_writes_sidecar(tmp_path):
    target = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("Sidecar body")
    document.save(target)
    result = {"path": str(target), "filename": "notes.docx"}
    _attach_extract(result, write_sidecar=True)
    assert result["extract_status"] == "ok"
    assert result["excerpt"]
    sidecar = Path(result["extracted_text_path"])
    assert sidecar.read_text(encoding="utf-8").startswith("Sidecar body")


def test_extract_command_json(tmp_path, capsys):
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("CLI extract works")
    document.save(path)
    code = main(["moodle", "extract", "--path", str(path), "--json"])
    import json

    payload = json.loads(capsys.readouterr().out)
    assert code == 0
    assert payload["ok"] is True
    assert "CLI extract works" in payload["text"]
    assert payload["source"] in {"moodle_html", "unverified"}
