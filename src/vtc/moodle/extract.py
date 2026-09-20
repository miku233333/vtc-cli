"""Local text extraction for Moodle downloads. Never log file contents here."""

from __future__ import annotations

import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from vtc.moodle.html import clean_text

EXCERPT_CHARS = 2000
DEFAULT_MAX_CHARS = 12000
MAX_PDF_PAGES = 20
MAX_PPTX_SLIDES = 30

SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md"}
LEGACY_OFFICE = {".doc", ".ppt", ".rtf"}


@dataclass
class ExtractedText:
    status: str
    extractor: str | None = None
    text: str | None = None
    excerpt: str | None = None
    key_lines: list[str] = field(default_factory=list)
    error: str | None = None


def extract_file(path: Path, *, max_chars: int = DEFAULT_MAX_CHARS) -> ExtractedText:
    target = Path(path)
    if not target.exists() or not target.is_file():
        return ExtractedText(status="unverified", error="file_missing")

    suffix = target.suffix.lower()
    try:
        if suffix == ".pdf":
            text, extractor = _extract_pdf(target, max_chars)
        elif suffix == ".docx":
            text, extractor = _extract_docx(target, max_chars)
        elif suffix == ".pptx":
            text, extractor = _extract_pptx(target, max_chars)
        elif suffix in {".txt", ".md"}:
            raw = target.read_text(encoding="utf-8", errors="ignore")
            text, extractor = raw[:max_chars], "plain-text"
        elif suffix in LEGACY_OFFICE:
            text, extractor = _extract_textutil(target, max_chars)
            if not text:
                return ExtractedText(
                    status="unverified",
                    extractor=None,
                    error="legacy_office_unreadable",
                )
        else:
            return ExtractedText(status="unsupported", error=f"unsupported_type:{suffix or 'none'}")
    except Exception:
        return ExtractedText(status="unverified", error="extract_failed")

    if not text or not clean_text(text):
        return ExtractedText(status="unverified", extractor=extractor, error="no_extractable_text")

    excerpt = text[:EXCERPT_CHARS]
    return ExtractedText(
        status="ok",
        extractor=extractor,
        text=text,
        excerpt=excerpt,
        key_lines=_key_lines(text),
    )


def _extract_pdf(path: Path, max_chars: int) -> tuple[str | None, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    chunks: list[str] = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        chunks.append(page.extract_text() or "")
    text = "\n".join(chunks).strip()
    return (text[:max_chars] if text else None), "pypdf"


def _extract_docx(path: Path, max_chars: int) -> tuple[str | None, str]:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if clean_text(paragraph.text)]
    for table in document.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    return (text[:max_chars] if text else None), "python-docx"


def _extract_pptx(path: Path, max_chars: int) -> tuple[str | None, str]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides):
        if index >= MAX_PPTX_SLIDES:
            break
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = clean_text(shape.text_frame.text)
                if text:
                    parts.append(text)
            elif hasattr(shape, "text"):
                text = clean_text(str(shape.text))
                if text:
                    parts.append(text)
    text = "\n".join(parts).strip()
    return (text[:max_chars] if text else None), "python-pptx"


def _extract_textutil(path: Path, max_chars: int) -> tuple[str | None, str | None]:
    binary = Path("/usr/bin/textutil")
    if not binary.exists():
        return None, None
    try:
        result = subprocess.run(
            [str(binary), "-convert", "txt", "-stdout", str(path)],
            capture_output=True,
            text=True,
            check=True,
            timeout=30,
        )
    except Exception:
        return None, None
    text = (result.stdout or "").strip()
    return (text[:max_chars] if text else None), "textutil"


def _key_lines(text: str, limit: int = 12) -> list[str]:
    selected: list[str] = []
    for raw in text.splitlines():
        line = clean_text(raw)
        if not line or len(line) < 4:
            continue
        if line in selected:
            continue
        selected.append(line)
        if len(selected) >= limit:
            break
    return selected
